# -*- coding: utf-8 -*-
import re, json, glob, csv, unicodedata
from openpyxl import load_workbook
E='C:/Users/CARLOS/elnino_ibk/'
def pad(x):
    d=re.sub(r'\D','',str(x or '')); return d.zfill(6) if d else ''
def norm(s):
    s=unicodedata.normalize('NFD',str(s or '')); s=''.join(c for c in s if unicodedata.category(c)!='Mn')
    s=s.upper(); s=re.sub(r'[^A-Z0-9 ]',' ',s); return re.sub(r'\s+',' ',s).strip()

# ---------- 1) impacto humano 2017 (SINPAD, filtro FEN identico al pipeline) ----------
INCL=re.compile(r'LLUVIA|INUNDAC|HUAIC|HUAYC|ALUVI|ANIEGO|TEMPORAL|EROSION FLUVIAL|AGENTE CLIM|EMBALSE|DESLIZ|DERRUMBE CERRO')
EXCL=re.compile(r'DERRUMBE DE ROCA|REPTAC|LAHAR|DERRUMBE VIVIENDA|COLAPSO POR|COLAPSO DE ESTRUCTURA POR (ACTIVIDAD|AGENTE BIOL)')
def isFEN(s):
    u=str(s or '').upper(); return bool(INCL.search(u)) and not EXCL.search(u)
wb=load_workbook(E+'01_inputs/sinpad/BD_emer_danos_2003_2020.xlsx',read_only=True)
ws=wb['BD 2003-2020']
imp={}
for r in ws.iter_rows(min_row=4,values_only=True):
    try:
        if int(r[2] or 0)!=2017 or not isFEN(r[8]): continue
    except Exception: continue
    u=pad(r[4])
    if not u: continue
    d=imp.setdefault(u,{'dam':0,'afe':0,'fall':0,'viv':0})
    d['dam']+=int(r[13] or 0); d['afe']+=int(r[14] or 0)
    d['fall']+=int(r[10] or 0); d['viv']+=int(r[15] or 0)+int(r[16] or 0)
print('distritos con impacto 2017 FEN:',len(imp))
json.dump(imp,open('impacto2017.json','w',encoding='utf-8'),ensure_ascii=False)

# ---------- 2) noticias/evidencia por distrito ----------
# mapa (prov|dist)->ubigeo via v6
p2u={}
with open(E+'02_pipeline/out/grado_fen_v6.csv',encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        p2u[norm(row['provincia'])+'|'+norm(row['distrito'])]=row['ubigeo']

news={}
def add(dep,prov,dist,estado,ev,url,grado=''):
    u=p2u.get(norm(prov)+'|'+norm(dist))
    if not u or not (ev or url): return
    lst=news.setdefault(u,[])
    if any(x.get('url')==url and url for x in lst): return
    if len(lst)>=2: return
    lst.append({'estado':(estado or '').strip(),'ev':(ev or '').strip(),'url':(url or '').strip(),'grado':grado})

# verificacion_r4r5.csv
with open(E+'07_validacion_noticias/verificacion_r4r5.csv',encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        add(row.get('departamento'),row.get('provincia'),row.get('distrito'),row.get('estado'),row.get('evidencia'),row.get('fuente'),row.get('grado'))
# batches json
for fn in sorted(glob.glob(E+'07_validacion_noticias/batch_*_resultado.json')):
    try: arr=json.load(open(fn,encoding='utf-8'))
    except Exception: continue
    for o in arr:
        add(o.get('departamento'),o.get('provincia'),o.get('distrito'),o.get('estado'),o.get('evidencia'),o.get('fuente'),o.get('grado'))
# spotcheck_r1r2.csv (si tiene columnas compatibles)
try:
    with open(E+'07_validacion_noticias/spotcheck_r1r2.csv',encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            add(row.get('departamento'),row.get('provincia'),row.get('distrito'),row.get('estado') or row.get('veredicto'),row.get('evidencia'),row.get('fuente'))
except Exception as e: print('spotcheck skip:',e)

print('distritos con noticias:',len(news))
json.dump(news,open('noticias.json','w',encoding='utf-8'),ensure_ascii=False)
# muestra
for u in list(news)[:3]:
    print(u,'->',news[u][0].get('estado'),news[u][0].get('url')[:70])
